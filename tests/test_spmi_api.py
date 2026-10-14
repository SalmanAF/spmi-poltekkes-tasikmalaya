import unittest
from io import BytesIO
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app, require_admin


class SpmiApiTests(unittest.TestCase):
    def setUp(self):
        self.documents = [
            {
                "id": "std-edu",
                "judul": "Standar Pendidikan",
                "kategori": "Standar PT",
                "kode": "STD-EDU",
                "tahun": "2025",
                "created_at": "2025-01-01",
                "status": "publik",
                "standard_id": "std-edu",
                "file_url": "",
                "tipe_sumber": "drive_link",
                "is_downloadable": False,
            },
            {
                "id": "guide-edu",
                "judul": "Pedoman Pendidikan",
                "kategori": "Pedoman",
                "kode": "PED-EDU",
                "tahun": "2025",
                "created_at": "2025-02-01",
                "status": "publik",
                "standard_id": "std-edu",
                "file_url": "",
                "tipe_sumber": "drive_link",
                "is_downloadable": False,
            },
            {
                "id": "sop-edu",
                "judul": "SOP Pendidikan",
                "kategori": "SOP",
                "kode": "SOP-EDU",
                "tahun": "2025",
                "created_at": "2025-03-01",
                "status": "publik",
                "standard_id": "std-edu",
                "guideline_id": "guide-edu",
                "parent_document_id": "guide-edu",
                "file_url": "",
                "tipe_sumber": "drive_link",
                "is_downloadable": False,
            },
            {
                "id": "draft-std",
                "judul": "Draft Standar",
                "kategori": "Standar PT",
                "kode": "DRAFT-STD",
                "tahun": "2025",
                "created_at": "2025-04-01",
                "status": "draft",
                "standard_id": "draft-std",
                "file_url": "",
                "tipe_sumber": "drive_link",
                "is_downloadable": False,
            },
            {
                "id": "std-2022",
                "judul": "Standar Tata Kelola 2022",
                "kategori": "Standar PT",
                "kode": "STD-2022",
                "tahun": "2026",
                "tahun_integrasi": "2022",
                "created_at": "2026-01-01",
                "status": "publik",
                "standard_id": "std-2022",
                "file_url": "",
                "tipe_sumber": "drive_link",
                "is_downloadable": False,
            },
        ]
        app.dependency_overrides[require_admin] = lambda: None
        self.patches = [
            patch("app.routers.spmi.all_documents", return_value=self.documents),
            patch(
                "app.routers.spmi._integration_records",
                return_value=self.documents[:3],
            ),
            patch("app.routers.spmi.implementation_overview", return_value=[]),
            patch(
                "app.routers.spmi.implementation_stage_counts",
                return_value={
                    "penetapan": 0,
                    "pelaksanaan": 0,
                    "evaluasi": 0,
                    "pengendalian": 0,
                    "peningkatan": 0,
                },
            ),
            patch(
                "app.routers.spmi.implementation_stage_complete_counts",
                return_value={
                    "penetapan": 0,
                    "pelaksanaan": 0,
                    "evaluasi": 0,
                    "pengendalian": 0,
                    "peningkatan": 0,
                },
            ),
        ]
        for mocked in self.patches:
            mocked.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.pop(require_admin, None)
        for mocked in reversed(self.patches):
            mocked.stop()

    def test_structure_summary_and_tree(self):
        summary = self.client.get("/api/v1/struktur-dokumen/summary")
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.json()["total_dokumen"], 5)
        self.assertEqual(summary.json()["total_publik"], 4)
        self.assertEqual(summary.json()["total_non_publik"], 1)
        self.assertEqual(summary.json()["per_tipe"]["Standar PT"], 3)

        tree = self.client.get("/api/v1/struktur-dokumen/tree")
        self.assertEqual(tree.status_code, 200)
        standard = tree.json()["standards"][0]
        self.assertEqual(standard["id"], "std-edu")
        self.assertEqual(standard["children"][0]["id"], "guide-edu")
        self.assertEqual(standard["children"][0]["children"][0]["id"], "sop-edu")

    def test_structure_can_be_filtered_by_integration_year(self):
        summary = self.client.get(
            "/api/v1/struktur-dokumen/summary",
            params={"tahun": "2022"},
        )
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.json()["filter_tahun"], "2022")
        self.assertEqual(summary.json()["per_tipe"]["Standar PT"], 1)
        self.assertIn("2022", summary.json()["tahun_tersedia"])

        tree = self.client.get(
            "/api/v1/struktur-dokumen/tree",
            params={"tahun": "2022"},
        )
        self.assertEqual(tree.status_code, 200)
        self.assertEqual(
            [standard["id"] for standard in tree.json()["standards"]],
            ["std-2022"],
        )

        search = self.client.get(
            "/api/v1/struktur-dokumen/search",
            params={"tahun": "2022"},
        )
        self.assertEqual(search.status_code, 200)
        self.assertEqual(
            [document["id"] for document in search.json()["documents"]],
            ["std-2022"],
        )

    def test_pillar_map_and_detail(self):
        result = self.client.get("/api/v1/peta-integrasi").json()
        education = next(item for item in result["pilar"] if item["pilar_id"] == "pendidikan")
        self.assertEqual(education["jumlah_standar"], 1)
        self.assertEqual(education["jumlah_pedoman"], 1)
        self.assertEqual(education["jumlah_sop"], 1)

        detail = self.client.get("/api/v1/peta-integrasi/pendidikan/detail")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()["total_dokumen"], 3)
        self.assertEqual(
            self.client.get("/api/v1/peta-integrasi/unknown/detail").status_code,
            404,
        )

    def test_admin_search_and_reports(self):
        search = self.client.get(
            "/api/v1/struktur-dokumen/search",
            params={"keyword": "Draft", "status": "draft"},
        )
        self.assertEqual(search.status_code, 200)
        self.assertEqual(search.json()["total"], 1)
        self.assertEqual(search.json()["documents"][0]["id"], "draft-std")

        completeness = self.client.get("/api/v1/laporan/kelengkapan")
        self.assertEqual(completeness.status_code, 200)
        self.assertEqual(completeness.json()["total_dokumen"], 5)
        self.assertFalse(completeness.json()["documents"][0]["siap_audit"])

        revision = self.client.get("/api/v1/laporan/revisi")
        self.assertFalse(revision.json()["data_tersedia"])
        self.assertEqual(revision.json()["total"], 0)

    def test_revision_report_uses_expiry_and_review_metadata(self):
        records = [
            {
                **self.documents[0],
                "berlaku_sampai": "2000-01-01",
                "status_revisi": "dalam_peninjauan",
                "sumber_revisi": "audit",
                "catatan_revisi": "Tindak lanjut temuan audit.",
                "ditinjau_pada": "2026-01-15",
            },
            self.documents[1],
        ]
        with patch("app.routers.spmi.all_documents", return_value=records):
            report = self.client.get("/api/v1/laporan/revisi").json()
        self.assertTrue(report["data_tersedia"])
        self.assertEqual(report["kedaluwarsa"], 1)
        self.assertEqual(report["dalam_peninjauan"], 1)
        self.assertEqual(report["total"], 1)
        self.assertIn("audit", report["documents"][0]["sumber_revisi"])

    def test_reports_include_completion_percent_and_orphans(self):
        completeness = self.client.get("/api/v1/laporan/kelengkapan").json()
        self.assertEqual(completeness["persentase_siap_audit"], 0)

        integration = self.client.get("/api/v1/laporan/integrasi").json()
        self.assertGreater(integration["integrasi_belum_lengkap"], 0)
        self.assertTrue(integration["dokumen_yatim"])
        self.assertEqual(
            integration["total_dokumen"],
            len([item for item in self.documents if item["kategori"] in {"Standar PT", "Pedoman", "SOP"}]),
        )

    def test_integration_report_rejects_mismatched_sop_standard(self):
        records = [dict(item) for item in self.documents[:3]]
        records[2]["standard_id"] = "std-2022"
        with patch("app.routers.spmi.all_documents", return_value=records):
            report = self.client.get("/api/v1/laporan/integrasi").json()
        orphan_ids = {item["id"] for item in report["dokumen_yatim"]}
        self.assertIn("sop-edu", orphan_ids)
        self.assertIn("guide-edu", orphan_ids)
        self.assertEqual(report["integrasi_lengkap"], 0)

    def test_reports_require_admin(self):
        app.dependency_overrides.pop(require_admin, None)
        response = self.client.get("/api/v1/laporan/kelengkapan")
        self.assertEqual(response.status_code, 401)
        self.assertEqual(
            self.client.get("/api/v1/struktur-dokumen/tree").status_code,
            401,
        )

    def test_excel_and_pdf_exports(self):
        with patch("app.routers.spmi.all_documents", return_value=[]):
            excel = self.client.get("/api/v1/laporan/export/excel")
            pdf = self.client.get("/api/v1/laporan/export/pdf")
        self.assertEqual(excel.status_code, 200)
        self.assertTrue(
            excel.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        self.assertTrue(excel.content.startswith(b"PK"))
        workbook = load_workbook(BytesIO(excel.content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Ringkasan",
                "Dokumen",
                "Kelengkapan",
                "Integrasi",
                "Relasi Belum Lengkap",
                "Implementasi",
                "Perkembangan",
                "Revisi",
            ],
        )
        workbook.close()
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_new_export_routes_and_print_page(self):
        with patch("app.routers.spmi.all_documents", return_value=[]):
            self.assertEqual(
                self.client.get("/api/v1/laporan/export-excel").status_code,
                200,
            )
            self.assertEqual(
                self.client.get("/api/v1/laporan/export-pdf").status_code,
                200,
            )
            printed = self.client.get("/api/v1/laporan/print")
        self.assertEqual(printed.status_code, 200)
        self.assertIn("text/html", printed.headers["content-type"])
        self.assertIn("window.print()", printed.text)


if __name__ == "__main__":
    unittest.main()
