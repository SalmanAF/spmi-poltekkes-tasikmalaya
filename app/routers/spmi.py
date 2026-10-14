from __future__ import annotations

from collections import Counter
from datetime import date
from html import escape
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.implementation import (
    PPEPP_STAGES,
    implementation_overview,
    implementation_stage_complete_counts,
    implementation_stage_counts,
)
from app.schemas import (
    CompletenessItem,
    CompletenessReport,
    DevelopmentReport,
    DocumentRelationship,
    DocumentSummary,
    DocumentTrend,
    ImplementationReport,
    IntegrationMapResponse,
    IntegrationPillar,
    IntegrationPillarDetail,
    IntegrationReport,
    IntegrationStandard,
    ReportExportSummary,
    RevisionItem,
    RevisionReport,
    StructureNode,
    StructureSearchResponse,
    StructureSummary,
    StructureTreeResponse,
)
from app.user_documents import (
    all_documents,
    is_google_drive_url,
    safe_local_path,
)


admin_router = APIRouter(prefix="/api/v1")

PILLARS: tuple[tuple[str, str], ...] = (
    ("pendidikan", "Standar Pendidikan"),
    ("penelitian", "Standar Penelitian"),
    ("pengabdian", "Standar Pengabdian"),
    ("tambahan", "Standar Tambahan"),
)
DOCUMENT_TYPES = {"Standar PT", "Pedoman", "SOP"}
INTEGRATION_YEARS = {"2022", "2025"}


def _integration_records() -> list[dict]:
    return [
        document
        for document in all_documents()
        if str(document.get("tahun_integrasi") or "") in INTEGRATION_YEARS
    ]


def _has_file(document: dict) -> bool:
    file_url = str(document.get("file_url") or "")
    if document.get("tipe_sumber") == "local_file":
        local_path = safe_local_path(file_url)
        return bool(local_path and local_path.is_file())
    return bool(file_url and is_google_drive_url(file_url))


def _summary(document: dict) -> DocumentSummary:
    return DocumentSummary(
        id=str(document.get("id") or ""),
        judul=str(document.get("judul") or ""),
        kategori=str(document.get("kategori") or ""),
        kode=str(document.get("kode") or document.get("no_dokumen") or "-"),
        status=document.get("status", "draft"),
        tahun=str(document.get("tahun") or ""),
        tahun_integrasi=str(document.get("tahun_integrasi") or "") or None,
        has_file=_has_file(document),
        download_allowed=bool(document.get("is_downloadable")),
        standard_id=str(document.get("standard_id") or ""),
        guideline_id=str(document.get("guideline_id") or ""),
        parent_document_id=str(document.get("parent_document_id") or ""),
    )


def _year_matches(document: dict, year: str | None) -> bool:
    return not year or str(
        document.get("tahun_integrasi") or document.get("tahun") or ""
    ) == year


def _record_by_id(records: list[dict]) -> dict[str, dict]:
    return {str(record.get("id") or ""): record for record in records}


def _build_tree(
    records: list[dict],
    year: str | None = None,
) -> StructureTreeResponse:
    by_id = _record_by_id(records)
    standards = [
        record
        for record in records
        if record.get("kategori") == "Standar PT"
        and _year_matches(record, year)
    ]
    guidelines = [
        record
        for record in records
        if record.get("kategori") == "Pedoman"
        and _year_matches(record, year)
    ]
    procedures = [
        record
        for record in records
        if record.get("kategori") == "SOP" and _year_matches(record, year)
    ]
    matching_ids = {
        str(record.get("id") or "")
        for record in records
        if record.get("kategori") in DOCUMENT_TYPES and _year_matches(record, year)
    }
    matching_guideline_ids = {
        str(record.get("guideline_id") or record.get("parent_document_id") or "")
        for record in procedures
    }
    matching_standard_ids = {
        str(record.get("standard_id") or "")
        for record in guidelines + procedures
    }
    standards_by_id = {
        str(record.get("id") or ""): record
        for record in records
        if record.get("kategori") == "Standar PT"
    }
    guidelines_by_id = {
        str(record.get("id") or ""): record
        for record in records
        if record.get("kategori") == "Pedoman"
    }
    for guideline_id in matching_guideline_ids:
        guideline = guidelines_by_id.get(guideline_id)
        if guideline:
            matching_standard_ids.add(str(guideline.get("standard_id") or ""))
    if year:
        standards = [
            record
            for record in standards_by_id.values()
            if str(record.get("id") or "") in matching_ids
            or str(record.get("id") or "") in matching_standard_ids
        ]
        guidelines = [
            record
            for record in guidelines_by_id.values()
            if str(record.get("id") or "") in matching_ids
            or str(record.get("id") or "") in matching_guideline_ids
        ]
    linked_ids: set[str] = set()
    roots: list[StructureNode] = []

    for standard in standards:
        standard_id = str(standard.get("id") or "")
        linked_guidelines = [
            record
            for record in guidelines
            if str(record.get("standard_id") or "") == standard_id
        ]
        guideline_nodes: list[StructureNode] = []
        for guideline in linked_guidelines:
            guideline_id = str(guideline.get("id") or "")
            linked_procedures = [
                record
                for record in procedures
                if str(record.get("guideline_id") or record.get("parent_document_id") or "")
                == guideline_id
            ]
            linked_ids.add(guideline_id)
            linked_ids.update(str(item.get("id") or "") for item in linked_procedures)
            guideline_nodes.append(
                StructureNode(
                    **_summary(guideline).model_dump(),
                    children=[
                        StructureNode(**_summary(item).model_dump())
                        for item in linked_procedures
                    ],
                )
            )

        direct_procedures = []
        for record in procedures:
            if str(record.get("standard_id") or "") != standard_id:
                continue
            parent_id = str(
                record.get("guideline_id") or record.get("parent_document_id") or ""
            )
            parent = by_id.get(parent_id)
            if not parent_id or not parent or parent.get("kategori") != "Pedoman":
                direct_procedures.append(record)
        linked_ids.add(standard_id)
        linked_ids.update(str(item.get("id") or "") for item in direct_procedures)
        roots.append(
            StructureNode(
                **_summary(standard).model_dump(),
                children=[
                    *guideline_nodes,
                    *[
                        StructureNode(**_summary(item).model_dump())
                        for item in direct_procedures
                    ],
                ],
            )
        )

    for guideline in guidelines:
        standard_id = str(guideline.get("standard_id") or "")
        if standard_id in by_id and by_id[standard_id].get("kategori") == "Standar PT":
            linked_ids.add(str(guideline.get("id") or ""))

    unlinked = [
        _summary(record)
        for record in records
        if record.get("kategori") in DOCUMENT_TYPES
        and str(record.get("id") or "") not in linked_ids
        and _year_matches(record, year)
    ]
    return StructureTreeResponse(
        standards=roots,
        unlinked=unlinked,
        total_standards=sum(
            record.get("kategori") == "Standar PT"
            and _year_matches(record, year)
            for record in records
        ),
    )


def _pillar_id(standard: dict) -> str:
    text = " ".join(
        str(standard.get(field) or "")
        for field in ("subjudul", "judul")
    ).casefold()
    if "pendidikan" in text or "pembelajaran" in text:
        return "pendidikan"
    if "penelitian" in text:
        return "penelitian"
    if "pengabdian" in text or "pkm" in text:
        return "pengabdian"
    return "tambahan"


def _pillar_data(records: list[dict]) -> list[IntegrationPillar]:
    standards = [record for record in records if record.get("kategori") == "Standar PT"]
    result: list[IntegrationPillar] = []
    for pillar_id, label in PILLARS:
        pillar_standards: list[IntegrationStandard] = []
        for standard in standards:
            if _pillar_id(standard) != pillar_id:
                continue
            standard_id = str(standard.get("id") or "")
            guidelines = [
                item
                for item in records
                if item.get("kategori") == "Pedoman"
                and str(item.get("standard_id") or "") == standard_id
            ]
            guideline_ids = {str(item.get("id") or "") for item in guidelines}
            procedures = [
                item
                for item in records
                if item.get("kategori") == "SOP"
                and (
                    str(item.get("guideline_id") or item.get("parent_document_id") or "")
                    in guideline_ids
                    or (
                        str(item.get("standard_id") or "") == standard_id
                        and not str(item.get("guideline_id") or item.get("parent_document_id") or "")
                    )
                )
            ]
            pillar_standards.append(
                IntegrationStandard(
                    standard=_summary(standard),
                    pedoman=[_summary(item) for item in guidelines],
                    sop=[_summary(item) for item in procedures],
                )
            )
        result.append(
            IntegrationPillar(
                pilar_id=pillar_id,
                nama=label,
                jumlah_standar=len(pillar_standards),
                jumlah_pedoman=sum(len(item.pedoman) for item in pillar_standards),
                jumlah_sop=sum(len(item.sop) for item in pillar_standards),
                standards=pillar_standards,
            )
        )
    return result


def _completeness_report(records: list[dict]) -> CompletenessReport:
    records = [record for record in records if record.get("kategori") in DOCUMENT_TYPES]
    criteria = [
        "Judul dokumen terisi",
        "Nomor/kode dokumen terisi",
        "Tahun dokumen terisi",
        "Berkas/link dokumen tersedia",
        "Status dokumen bukan draft",
        "Pedoman memiliki relasi ke Standar PT yang valid",
        "SOP memiliki relasi ke Pedoman yang valid",
    ]
    by_id = _record_by_id(records)
    items: list[CompletenessItem] = []
    for document in records:
        missing: list[str] = []
        if not str(document.get("judul") or "").strip():
            missing.append("judul")
        if str(document.get("kode") or document.get("no_dokumen") or "-").strip() in {"", "-"}:
            missing.append("kode")
        if not str(document.get("tahun_integrasi") or document.get("tahun") or "").strip():
            missing.append("tahun")
        if not _has_file(document):
            missing.append("file dokumen tersedia")
        if document.get("status") == "draft":
            missing.append("dokumen masih berstatus draft")
        category = document.get("kategori")
        if category == "Pedoman":
            parent = by_id.get(str(document.get("standard_id") or ""))
            if not parent or parent.get("kategori") != "Standar PT":
                missing.append("relasi Pedoman ke Standar PT")
        elif category == "SOP":
            parent_id = str(
                document.get("guideline_id")
                or document.get("parent_document_id")
                or ""
            )
            parent = by_id.get(parent_id)
            if not parent or parent.get("kategori") != "Pedoman":
                missing.append("relasi SOP ke Pedoman")
        items.append(
            CompletenessItem(
                document=_summary(document),
                siap_audit=not missing,
                kekurangan=missing,
            )
        )
    ready = sum(item.siap_audit for item in items)
    return CompletenessReport(
        total_dokumen=len(items),
        siap_audit=ready,
        belum_lengkap=len(items) - ready,
        persentase_siap_audit=round(ready * 100 / len(items), 2) if items else 0,
        kriteria=criteria,
        documents=items,
    )


def _integration_report(records: list[dict]) -> IntegrationReport:
    records = [record for record in records if record.get("kategori") in DOCUMENT_TYPES]
    by_id = _record_by_id(records)
    relationships: dict[tuple[str, str, str], DocumentRelationship] = {}
    counts: Counter[str] = Counter()
    guidelines_by_standard: dict[str, list[dict]] = {}
    procedures_by_guideline: dict[str, list[dict]] = {}
    for document in records:
        if document.get("kategori") == "Pedoman":
            guidelines_by_standard.setdefault(
                str(document.get("standard_id") or ""), []
            ).append(document)
        elif document.get("kategori") == "SOP":
            parent_id = str(
                document.get("guideline_id")
                or document.get("parent_document_id")
                or ""
            )
            procedures_by_guideline.setdefault(parent_id, []).append(document)

    valid_procedures_by_guideline: dict[str, list[dict]] = {}
    for guideline_id, procedures in procedures_by_guideline.items():
        guideline = by_id.get(guideline_id)
        standard_id = str((guideline or {}).get("standard_id") or "")
        standard = by_id.get(standard_id)
        if not guideline or guideline.get("kategori") != "Pedoman":
            continue
        if not standard or standard.get("kategori") != "Standar PT":
            continue
        valid_procedures_by_guideline[guideline_id] = [
            procedure
            for procedure in procedures
            if not procedure.get("standard_id")
            or str(procedure.get("standard_id")) == standard_id
        ]

    for document in records:
        document_id = str(document.get("id") or "")
        category = str(document.get("kategori") or "")
        parent_id = ""
        relationship_type = ""
        if category == "Pedoman":
            parent_id = str(document.get("standard_id") or "")
            relationship_type = "standar_ke_pedoman"
        elif category == "SOP":
            parent_id = str(
                document.get("guideline_id")
                or document.get("parent_document_id")
                or document.get("standard_id")
                or ""
            )
            if parent_id:
                parent_type = str(by_id.get(parent_id, {}).get("kategori") or "")
                relationship_type = (
                    "pedoman_ke_sop" if parent_type == "Pedoman" else "standar_ke_sop"
                )
        expected_parent_type = (
            "Standar PT" if category == "Pedoman" else "Pedoman"
        )
        if (
            parent_id
            and parent_id in by_id
            and by_id[parent_id].get("kategori") == expected_parent_type
            and document_id
        ):
            key = (parent_id, document_id, relationship_type)
            relationships[key] = DocumentRelationship(
                sumber_id=parent_id,
                sumber_tipe=str(by_id[parent_id].get("kategori") or ""),
                target_id=document_id,
                target_tipe=category,
                relasi=relationship_type,
            )
            counts[relationship_type] += 1

    standards = [item for item in records if item.get("kategori") == "Standar PT"]
    guidelines = [item for item in records if item.get("kategori") == "Pedoman"]
    procedures = [item for item in records if item.get("kategori") == "SOP"]
    standards_without_guidelines = [
        standard
        for standard in standards
        if not guidelines_by_standard.get(str(standard.get("id") or ""))
    ]
    guidelines_without_procedures = [
        guideline
        for guideline in guidelines
        if not valid_procedures_by_guideline.get(str(guideline.get("id") or ""))
    ]
    orphan_ids: set[str] = set()
    complete_ids: set[str] = set()

    for standard in standards:
        standard_id = str(standard.get("id") or "")
        related_guidelines = guidelines_by_standard.get(standard_id, [])
        if related_guidelines and all(
            valid_procedures_by_guideline.get(str(item.get("id") or ""))
            for item in related_guidelines
        ):
            complete_ids.add(standard_id)
        else:
            orphan_ids.add(standard_id)
    for guideline in guidelines:
        guideline_id = str(guideline.get("id") or "")
        standard = by_id.get(str(guideline.get("standard_id") or ""))
        has_valid_parent = bool(standard and standard.get("kategori") == "Standar PT")
        has_child = bool(valid_procedures_by_guideline.get(guideline_id))
        if has_valid_parent and has_child:
            complete_ids.add(guideline_id)
        else:
            orphan_ids.add(guideline_id)
    for procedure in procedures:
        procedure_id = str(procedure.get("id") or "")
        guideline = by_id.get(
            str(
                procedure.get("guideline_id")
                or procedure.get("parent_document_id")
                or ""
            )
        )
        parent_standard_id = str((guideline or {}).get("standard_id") or "")
        parent_standard = by_id.get(parent_standard_id)
        standard_matches = (
            bool(parent_standard and parent_standard.get("kategori") == "Standar PT")
            and (
                not procedure.get("standard_id")
                or str(procedure.get("standard_id")) == parent_standard_id
            )
        )
        if (
            guideline
            and guideline.get("kategori") == "Pedoman"
            and standard_matches
        ):
            complete_ids.add(procedure_id)
        else:
            orphan_ids.add(procedure_id)

    values = list(relationships.values())
    total = len(records)
    complete_count = len(complete_ids)
    return IntegrationReport(
        total_dokumen=total,
        total_relasi=len(values),
        integrasi_lengkap=complete_count,
        integrasi_belum_lengkap=total - complete_count,
        persentase_integrasi_lengkap=(
            round(complete_count * 100 / total, 2) if total else 0
        ),
        dokumen_yatim=[
            _summary(item)
            for item in records
            if str(item.get("id") or "") in orphan_ids
        ],
        standar_tanpa_pedoman=[_summary(item) for item in standards_without_guidelines],
        pedoman_tanpa_sop=[_summary(item) for item in guidelines_without_procedures],
        per_tipe_relasi=dict(counts),
        relationships=values,
    )


def _implementation_report(records: list[dict]) -> ImplementationReport:
    records = [record for record in records if record.get("kategori") in DOCUMENT_TYPES]
    overview = implementation_overview(records)
    tracked = sum(
        bool(str(item.get("unit") or "").strip())
        and bool(str(item.get("evidence") or "").strip())
        and bool(item.get("implementation_date"))
        for item in overview
    )
    stage_counts = implementation_stage_counts(records)
    complete_stage_counts = implementation_stage_complete_counts(records)
    missing_ppepp = [
        _summary(item)
        for item in overview
        if not item.get("ppepp_complete")
    ]
    return ImplementationReport(
        total_dokumen=len(overview),
        terpantau=tracked,
        belum_terpantau=len(overview) - tracked,
        memiliki_data_ppepp=sum(bool(item.get("has_ppepp")) for item in overview),
        ppepp_lengkap=sum(bool(item.get("ppepp_complete")) for item in overview),
        persentase_ppepp_lengkap=(
            round(
                sum(bool(item.get("ppepp_complete")) for item in overview)
                * 100
                / len(overview),
                2,
            )
            if overview
            else 0
        ),
        ppepp_belum_lengkap=missing_ppepp,
        per_tahap=stage_counts,
        tahap_lengkap=complete_stage_counts,
    )


def _development_report(records: list[dict]) -> DevelopmentReport:
    records = [record for record in records if record.get("kategori") in DOCUMENT_TYPES]
    counts: Counter[str] = Counter()
    for document in records:
        created_at = str(document.get("created_at") or "")
        year = created_at[:4]
        if len(year) == 4 and year.isdigit():
            counts[year] += 1
    cumulative = 0
    trends: list[DocumentTrend] = []
    for year, count in sorted(counts.items()):
        cumulative += count
        trends.append(
            DocumentTrend(tahun=year, dokumen_baru=count, kumulatif=cumulative)
        )
    return DevelopmentReport(
        total_dokumen=len(records),
        basis_tanggal="created_at",
        catatan=(
            "Tren dihitung dari tanggal pencatatan created_at. Untuk dokumen hasil "
            "impor, tanggal ini dapat merepresentasikan waktu impor, bukan tanggal "
            "pengesahan atau perubahan historis."
        ),
        perkembangan=trends,
    )


def _revision_report(records: list[dict]) -> RevisionReport:
    records = [record for record in records if record.get("kategori") in DOCUMENT_TYPES]
    today = date.today()
    items: list[RevisionItem] = []
    expired_count = 0
    review_count = 0
    revise_count = 0
    metadata_available = any(
        record.get("berlaku_sampai")
        or record.get("status_revisi")
        or record.get("sumber_revisi")
        or record.get("catatan_revisi")
        or record.get("ditinjau_pada")
        for record in records
    )
    for record in records:
        expiry = str(record.get("berlaku_sampai") or "").strip()
        review_status = str(record.get("status_revisi") or "").strip()
        reasons: list[str] = []
        is_expired = False
        if expiry:
            try:
                is_expired = date.fromisoformat(expiry) < today
            except ValueError as exc:
                raise RuntimeError(
                    f"Tanggal berlaku dokumen {record.get('id')} tidak valid."
                ) from exc
        if is_expired:
            expired_count += 1
            reasons.append("Masa berlaku dokumen telah berakhir")
        if review_status == "dalam_peninjauan":
            review_count += 1
            reasons.append("Dokumen sedang dalam peninjauan")
        elif review_status == "perlu_revisi":
            revise_count += 1
            reasons.append("Dokumen memerlukan revisi")
        if not reasons:
            continue
        items.append(
            RevisionItem(
                document=_summary(record),
                alasan=reasons,
                berlaku_sampai=expiry or None,
                status_revisi=review_status or None,
                sumber_revisi=record.get("sumber_revisi"),
                catatan_revisi=str(record.get("catatan_revisi") or ""),
                ditinjau_pada=str(record.get("ditinjau_pada") or ""),
            )
        )
    return RevisionReport(
        data_tersedia=metadata_available,
        total=len(items),
        message=(
            "Data revisi berdasarkan metadata masa berlaku dan catatan peninjauan."
            if metadata_available
            else (
                "Belum ada metadata masa berlaku, status revisi, atau catatan audit/"
                "masukan. Isi metadata melalui API pembaruan dokumen agar laporan "
                "revisi dapat dihitung."
            )
        ),
        kedaluwarsa=expired_count,
        dalam_peninjauan=review_count,
        perlu_revisi=revise_count,
        documents=items,
    )


def _report_summary(records: list[dict]) -> ReportExportSummary:
    completeness = _completeness_report(records)
    integration = _integration_report(records)
    implementation = _implementation_report(records)
    revision = _revision_report(records)
    return ReportExportSummary(
        total_dokumen=_completeness_report(records).total_dokumen,
        siap_audit=completeness.siap_audit,
        belum_lengkap=completeness.belum_lengkap,
        total_relasi=integration.total_relasi,
        terpantau=implementation.terpantau,
        ppepp_lengkap=implementation.ppepp_lengkap,
        persentase_siap_audit=completeness.persentase_siap_audit,
        persentase_integrasi_lengkap=integration.persentase_integrasi_lengkap,
        persentase_ppepp_lengkap=implementation.persentase_ppepp_lengkap,
        kedaluwarsa=revision.kedaluwarsa,
        dalam_peninjauan=revision.dalam_peninjauan,
        perlu_revisi=revision.perlu_revisi,
    )


@admin_router.get(
    "/struktur-dokumen/summary",
    response_model=StructureSummary,
)
async def struktur_dokumen_summary(
    tahun: str | None = Query(default=None, pattern=r"^\d{4}$"),
):
    documents = all_documents()
    years = sorted(
        {
            str(document.get("tahun_integrasi") or document.get("tahun") or "")
            for document in documents
            if str(document.get("tahun_integrasi") or document.get("tahun") or "").isdigit()
        },
        reverse=True,
    )
    filtered_documents = [
        document for document in documents if _year_matches(document, tahun)
    ]
    counts = Counter(
        str(document.get("kategori") or "")
        for document in filtered_documents
        if document.get("kategori") in DOCUMENT_TYPES
    )
    files_available = sum(_has_file(document) for document in filtered_documents)
    public_count = sum(
        document.get("status") == "publik" for document in filtered_documents
    )
    return StructureSummary(
        total_dokumen=len(filtered_documents),
        total_publik=public_count,
        total_non_publik=len(filtered_documents) - public_count,
        filter_tahun=tahun,
        tahun_tersedia=years,
        per_tipe={
            "Standar PT": counts["Standar PT"],
            "Pedoman": counts["Pedoman"],
            "SOP": counts["SOP"],
        },
        ketersediaan_file={
            "tersedia": files_available,
            "tidak_tersedia": len(filtered_documents) - files_available,
        },
    )


@admin_router.get(
    "/struktur-dokumen/tree",
    response_model=StructureTreeResponse,
)
async def struktur_dokumen_tree(
    tahun: str | None = Query(default=None, pattern=r"^\d{4}$"),
):
    return _build_tree(all_documents(), tahun)


@admin_router.get(
    "/struktur-dokumen/search",
    response_model=StructureSearchResponse,
)
async def struktur_dokumen_search(
    keyword: str = Query(default="", max_length=200),
    tahun: str | None = Query(default=None, pattern=r"^\d{4}$"),
    tipe: Literal["Standar PT", "Pedoman", "SOP"] | None = None,
    status: Literal["publik", "privat", "draft"] | None = None,
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
):
    documents = [
        document for document in all_documents() if _year_matches(document, tahun)
    ]
    if keyword.strip():
        needle = keyword.casefold().strip()
        documents = [
            document
            for document in documents
            if needle
            in " ".join(
                str(document.get(field) or "")
                for field in ("judul", "kode", "no_dokumen", "kategori")
            ).casefold()
        ]
    if tipe:
        documents = [item for item in documents if item.get("kategori") == tipe]
    if status:
        documents = [item for item in documents if item.get("status") == status]
    total = len(documents)
    return StructureSearchResponse(
        total=total,
        limit=limit,
        offset=offset,
        documents=[_summary(item) for item in documents[offset : offset + limit]],
    )


@admin_router.get(
    "/peta-integrasi",
    response_model=IntegrationMapResponse,
)
async def peta_integrasi():
    return IntegrationMapResponse(pilar=_pillar_data(_integration_records()))


@admin_router.get(
    "/peta-integrasi/{pilar_id}/detail",
    response_model=IntegrationPillarDetail,
)
async def peta_integrasi_detail(
    pilar_id: str,
):
    pillars = _pillar_data(_integration_records())
    pillar = next((item for item in pillars if item.pilar_id == pilar_id), None)
    if pillar is None:
        raise HTTPException(status_code=404, detail="Pilar integrasi tidak ditemukan.")
    documents = [
        document
        for standard in pillar.standards
        for document in [standard.standard, *standard.pedoman, *standard.sop]
    ]
    return IntegrationPillarDetail(
        pilar=pillar,
        total_dokumen=len(documents),
        documents=documents,
    )


@admin_router.get(
    "/laporan/kelengkapan",
    response_model=CompletenessReport,
)
async def laporan_kelengkapan():
    return _completeness_report(all_documents())


@admin_router.get(
    "/laporan/integrasi",
    response_model=IntegrationReport,
)
async def laporan_integrasi():
    return _integration_report(all_documents())


@admin_router.get(
    "/laporan/implementasi",
    response_model=ImplementationReport,
)
async def laporan_implementasi():
    return _implementation_report(all_documents())


@admin_router.get(
    "/laporan/revisi",
    response_model=RevisionReport,
)
async def laporan_revisi():
    return _revision_report(all_documents())


@admin_router.get(
    "/laporan/perkembangan",
    response_model=DevelopmentReport,
)
async def laporan_perkembangan():
    return _development_report(all_documents())


@admin_router.get("/laporan/export-excel")
@admin_router.get("/laporan/export/excel", include_in_schema=False)
async def laporan_export_excel():
    records = all_documents()
    summary = _report_summary(records)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Ringkasan"
    overview.append(["Indikator", "Jumlah"])
    for key, value in summary.model_dump().items():
        overview.append([key, value])
    documents_sheet = workbook.create_sheet("Dokumen")
    documents_sheet.append(["ID", "Judul", "Kategori", "Kode", "Status", "Tahun", "File tersedia"])
    for document in records:
        item = _summary(document)
        documents_sheet.append(
            [
                item.id,
                item.judul,
                item.kategori,
                item.kode,
                item.status,
                item.tahun,
                "Ya" if item.has_file else "Tidak",
            ]
        )
    completeness_sheet = workbook.create_sheet("Kelengkapan")
    completeness = _completeness_report(records)
    completeness_sheet.append(["ID", "Judul", "Kategori", "Siap audit", "Kekurangan"])
    for item in completeness.documents:
        completeness_sheet.append(
            [
                item.document.id,
                item.document.judul,
                item.document.kategori,
                "Ya" if item.siap_audit else "Tidak",
                "; ".join(item.kekurangan),
            ]
        )
    relationships_sheet = workbook.create_sheet("Integrasi")
    relationships_sheet.append(["Sumber ID", "Tipe sumber", "Target ID", "Tipe target", "Relasi"])
    integration = _integration_report(records)
    for relation in integration.relationships:
        relationships_sheet.append(list(relation.model_dump().values()))
    orphan_sheet = workbook.create_sheet("Relasi Belum Lengkap")
    orphan_sheet.append(["ID", "Judul", "Kategori", "Status"])
    orphan_ids = {
        item.id: item for item in integration.dokumen_yatim
    }
    for item in integration.standar_tanpa_pedoman:
        orphan_ids[item.id] = item
    for item in integration.pedoman_tanpa_sop:
        orphan_ids[item.id] = item
    for item in orphan_ids.values():
        orphan_sheet.append([item.id, item.judul, item.kategori, "Periksa relasi"])
    implementation_sheet = workbook.create_sheet("Implementasi")
    implementation = _implementation_report(records)
    implementation_sheet.append(["Indikator", "Jumlah"])
    for key, value in implementation.model_dump().items():
        if key in {"per_tahap", "tahap_lengkap"}:
            for stage, count in value.items():
                label = "PPEPP lengkap" if key == "tahap_lengkap" else "PPEPP tercatat"
                implementation_sheet.append([f"{label} · {stage}", count])
        elif key == "ppepp_belum_lengkap":
            for item in value:
                implementation_sheet.append(
                    [f"PPEPP belum lengkap · {item['id']}", item["judul"]]
                )
        else:
            implementation_sheet.append([key, value])
    development_sheet = workbook.create_sheet("Perkembangan")
    development_sheet.append(["Tahun", "Dokumen baru", "Kumulatif"])
    for item in _development_report(records).perkembangan:
        development_sheet.append([item.tahun, item.dokumen_baru, item.kumulatif])
    revision_sheet = workbook.create_sheet("Revisi")
    revision_report = _revision_report(records)
    revision_sheet.append(
        ["ID", "Judul", "Alasan", "Berlaku sampai", "Status revisi", "Sumber", "Catatan"]
    )
    if revision_report.documents:
        for item in revision_report.documents:
            revision_sheet.append(
                [
                    item.document.id,
                    item.document.judul,
                    "; ".join(item.alasan),
                    item.berlaku_sampai or "",
                    item.status_revisi or "",
                    item.sumber_revisi or "",
                    item.catatan_revisi,
                ]
            )
    else:
        revision_sheet.append(["", revision_report.message])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B7890")
        for column in sheet.columns:
            width = min(
                60,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
            sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="laporan-spmi.xlsx"'},
    )


@admin_router.get("/laporan/export-pdf")
@admin_router.get("/laporan/export/pdf", include_in_schema=False)
async def laporan_export_pdf():
    records = all_documents()
    summary = _report_summary(records)
    completeness = _completeness_report(records)
    integration = _integration_report(records)
    implementation = _implementation_report(records)
    revision = _revision_report(records)
    development = _development_report(records)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.alignment = TA_CENTER
    content = [
        Paragraph("Laporan Rekapitulasi SPMI", title_style),
        Paragraph(f"Tanggal laporan: {date.today().isoformat()}", styles["Normal"]),
        Spacer(1, 8 * mm),
    ]
    table_data = [["Indikator", "Jumlah"]]
    table_data.extend(
        [label.replace("_", " ").title(), value]
        for label, value in summary.model_dump().items()
    )
    table = Table(table_data, colWidths=[100 * mm, 35 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B7890")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9DCE5")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F7FA")]),
                ("PADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    content.append(table)
    sections = [
        (
            "Kelengkapan Dokumen",
            [
                ["Siap audit", "Belum lengkap", "Persentase siap audit"],
                [
                    completeness.siap_audit,
                    completeness.belum_lengkap,
                    f"{completeness.persentase_siap_audit:.2f}%",
                ],
            ],
        ),
        (
            "Integrasi Dokumen",
            [
                ["Terintegrasi lengkap", "Belum lengkap", "Persentase"],
                [
                    integration.integrasi_lengkap,
                    integration.integrasi_belum_lengkap,
                    f"{integration.persentase_integrasi_lengkap:.2f}%",
                ],
            ],
        ),
        (
            "Implementasi PPEPP",
            [
                ["Terpantau", "PPEPP lengkap", "PPEPP belum lengkap"],
                [
                    implementation.terpantau,
                    implementation.ppepp_lengkap,
                    len(implementation.ppepp_belum_lengkap),
                ],
            ],
        ),
        (
            "Revisi Dokumen",
            [
                ["Kedaluwarsa", "Dalam peninjauan", "Memerlukan revisi"],
                [revision.kedaluwarsa, revision.dalam_peninjauan, revision.perlu_revisi],
            ],
        ),
        (
            "Perkembangan Dokumen",
            [
                ["Tahun", "Dokumen dicatat", "Kumulatif"],
                *[
                    [item.tahun, item.dokumen_baru, item.kumulatif]
                    for item in development.perkembangan
                ],
            ],
        ),
    ]
    for heading, data in sections:
        content.extend([Spacer(1, 7 * mm), Paragraph(escape(heading), styles["Heading2"])])
        report_table = Table(data, repeatRows=1)
        report_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B7890")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9DCE5")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F7FA")]),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        content.append(report_table)
    content.extend(
        [
            Spacer(1, 7 * mm),
            Paragraph("Catatan metode", styles["Heading2"]),
            Paragraph(escape(development.catatan), styles["Normal"]),
            Paragraph(escape(revision.message), styles["Normal"]),
        ]
    )
    document.build(content)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="laporan-spmi.pdf"'},
    )


def _print_table(headers: list[str], rows: list[list[object]]) -> str:
    header_html = "".join(f"<th>{escape(str(value))}</th>" for value in headers)
    rows_html = "".join(
        "<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{header_html}</tr></thead><tbody>{rows_html}</tbody></table>"


@admin_router.get("/laporan/print", response_class=HTMLResponse)
async def laporan_print():
    records = all_documents()
    summary = _report_summary(records)
    completeness = _completeness_report(records)
    integration = _integration_report(records)
    implementation = _implementation_report(records)
    revision = _revision_report(records)
    development = _development_report(records)
    summary_rows = [
        [key.replace("_", " ").title(), value]
        for key, value in summary.model_dump().items()
    ]
    incomplete_rows = [
        [
            item.document.id,
            item.document.judul,
            item.document.kategori,
            ", ".join(item.kekurangan),
        ]
        for item in completeness.documents
        if not item.siap_audit
    ][:100]
    orphan_rows = [
        [item.id, item.judul, item.kategori]
        for item in integration.dokumen_yatim
    ][:100]
    revision_rows = [
        [
            item.document.id,
            item.document.judul,
            ", ".join(item.alasan),
            item.berlaku_sampai or "-",
            item.catatan_revisi or "-",
        ]
        for item in revision.documents
    ][:100]
    trend_rows = [
        [item.tahun, item.dokumen_baru, item.kumulatif]
        for item in development.perkembangan
    ]
    stage_rows = [
        [
            stage["label"],
            implementation.per_tahap.get(stage["key"], 0),
            implementation.tahap_lengkap.get(stage["key"], 0),
        ]
        for stage in PPEPP_STAGES
    ]
    body = f"""
<!doctype html>
<html lang="id">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Laporan Rekapitulasi SPMI</title>
  <style>
    body {{ color:#172b3d; font:14px/1.5 Arial,sans-serif; margin:28px auto; max-width:1100px; }}
    h1,h2 {{ color:#0b2942; }} h1 {{ text-align:center; }} .meta {{ text-align:center; color:#627487; }}
    table {{ border-collapse:collapse; margin:12px 0 24px; width:100%; }}
    th,td {{ border:1px solid #d5e1e8; padding:7px 9px; text-align:left; vertical-align:top; }}
    th {{ background:#eaf4f8; }} .note {{ color:#526477; font-size:12px; }}
    @media print {{ body {{ margin:12mm; max-width:none; }} .no-print {{ display:none; }} thead {{ display:table-header-group; }} tr {{ break-inside:avoid; }} }}
  </style>
</head>
<body>
  <div class="no-print"><button onclick="window.print()">Cetak laporan</button></div>
  <h1>Laporan Rekapitulasi SPMI</h1>
  <p class="meta">Tanggal laporan: {date.today().isoformat()}</p>
  <h2>Ringkasan Eksekutif</h2>
  {_print_table(["Indikator", "Nilai"], summary_rows)}
  <h2>Kelengkapan Dokumen</h2>
  <p>Siap audit: {completeness.siap_audit} dari {completeness.total_dokumen} ({completeness.persentase_siap_audit:.2f}%).</p>
  {_print_table(["ID", "Judul", "Kategori", "Kekurangan"], incomplete_rows)}
  <h2>Integrasi</h2>
  <p>Dokumen terintegrasi lengkap: {integration.integrasi_lengkap} dari {integration.total_dokumen} ({integration.persentase_integrasi_lengkap:.2f}%).</p>
  {_print_table(["ID", "Judul", "Kategori"], orphan_rows)}
  <h2>Implementasi PPEPP</h2>
  <p>Terpantau: {implementation.terpantau}; PPEPP lengkap: {implementation.ppepp_lengkap} ({implementation.persentase_ppepp_lengkap:.2f}%).</p>
  {_print_table(["Tahap", "Tercatat", "Lengkap"], stage_rows)}
  <h2>Revisi Dokumen</h2>
  <p>{escape(revision.message)}</p>
  {_print_table(["ID", "Judul", "Alasan", "Berlaku sampai", "Catatan"], revision_rows)}
  <h2>Perkembangan Dokumen</h2>
  <p class="note">{escape(development.catatan)}</p>
  {_print_table(["Tahun pencatatan", "Dokumen baru", "Kumulatif"], trend_rows)}
</body>
</html>
"""
    return HTMLResponse(body)
